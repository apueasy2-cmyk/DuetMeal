import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = docx.Document()
    
    # Page Margins (1 inch everywhere)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header.is_linked_to_previous = False
        section.different_first_page_header_footer = False

    # Set normal style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    # Helper to format headings
    def add_sec_heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        if level == 1:
            run.font.size = Pt(13)
        elif level == 2:
            run.font.size = Pt(12)
        elif level == 3:
            run.font.size = Pt(11.5)
        return p

    def add_body_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        return p

    def add_bullet_p(title, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run_title = p.add_run(title + (": " if title else ""))
        run_title.bold = True
        run_title.font.name = 'Times New Roman'
        run_title.font.size = Pt(11)
        
        run_text = p.add_run(text)
        run_text.font.name = 'Times New Roman'
        run_text.font.size = Pt(11)
        return p

    def add_screenshot_box(fig_caption, height_in_inches=2.2):
        # 1-cell bordered table representing the blank image placeholder
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(5.8)
        
        cell = tbl.cell(0, 0)
        cell.width = Inches(5.8)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
        # Border and light shading for cell
        tcPr = cell._element.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="dashed" w:sz="6" w:space="0" w:color="A0AEC0"/>\n'
            f'  <w:left w:val="dashed" w:sz="6" w:space="0" w:color="A0AEC0"/>\n'
            f'  <w:bottom w:val="dashed" w:sz="6" w:space="0" w:color="A0AEC0"/>\n'
            f'  <w:right w:val="dashed" w:sz="6" w:space="0" w:color="A0AEC0"/>\n'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)
        
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
        tcPr.append(shading)
        
        # Inside placeholder text
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(int(height_in_inches * 26))
        cp.paragraph_format.space_after = Pt(int(height_in_inches * 26))
        crun = cp.add_run("[ INSERT SCREENSHOT HERE ]")
        crun.font.name = 'Times New Roman'
        crun.font.size = Pt(10.5)
        crun.font.italic = True
        crun.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
        
        # Figure caption paragraph below placeholder
        cap_p = doc.add_paragraph()
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_p.paragraph_format.space_before = Pt(4)
        cap_p.paragraph_format.space_after = Pt(12)
        cap_run = cap_p.add_run(fig_caption)
        cap_run.font.name = 'Times New Roman'
        cap_run.font.size = Pt(10.5)
        cap_run.bold = False

    # ==========================================
    # PAGE 1: INDEX
    # ==========================================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(20)
    r_title = p_title.add_run("Index")
    r_title.bold = True
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(16)

    # Sub-header
    p_head = doc.add_paragraph()
    p_head.paragraph_format.space_after = Pt(10)
    r_h1 = p_head.add_run("Report Content")
    r_h1.bold = True
    r_h1.font.size = Pt(11)
    
    # 100% exact match between Index headings and document body headings
    index_items = [
        ("1. Abstract", "2"),
        ("2. Introduction", "2"),
        ("3. Objectives", "2"),
        ("4. Key Features", "3"),
        ("5. Methodology", "3"),
        ("6. Project Overview", "4"),
        ("    6.0 User Login and Sign Up Section", "4"),
        ("        6.0.1 User Sign In Section", "4"),
        ("        6.0.2 User Sign Up Section", "4"),
        ("    6.1 Home Dashboard & Daily Meal Status", "5"),
        ("    6.2 Meal Booking & Date-Range Selection", "5"),
        ("    6.3 Check Availability & Guest Meal Reservation", "5"),
        ("    6.4 Review and Manage Your Canteen Wallet", "6"),
        ("    6.5 Daily Menu & Cutoff Schedule", "6"),
        ("    6.6 My Order & Meal History", "7"),
        ("    6.7 My Profile Section", "7"),
        ("    6.8 FAQ & Notice Board Section", "8"),
        ("7. Admin Panel", "9"),
        ("    7.1 Admin Sign-In Section", "9"),
        ("    7.2 Admin Dashboard Section", "10"),
        ("    7.3 Admin User Management Section", "10"),
        ("    7.4 Admin Menu Management Section", "11"),
        ("    7.5 Admin Recharge Approvals & Transactions", "11"),
        ("    7.6 Admin Monthly Settlement Section", "12"),
        ("8. Footer", "12"),
        ("    8.1 Quick Links, Help, and Contact in Footer Section", "12"),
        ("9. Challenges", "13"),
        ("10. Limitations", "13"),
        ("11. Learning Outcomes", "13"),
        ("12. Conclusion", "13"),
    ]

    idx_tbl = doc.add_table(rows=0, cols=2)
    idx_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    idx_tbl.autofit = False
    idx_tbl.columns[0].width = Inches(5.8)
    idx_tbl.columns[1].width = Inches(0.5)

    for item, pg in index_items:
        row = idx_tbl.add_row()
        c0 = row.cells[0]
        c1 = row.cells[1]
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        p0.paragraph_format.line_spacing = 1.05
        
        # Calculate dots
        dot_count = max(5, 75 - len(item) * 1)
        dots = " ." * (dot_count // 2)
        
        r0 = p0.add_run(f"{item} {dots}")
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(10)
        if not item.startswith(" "):
            r0.bold = True

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        p1.paragraph_format.line_spacing = 1.05
        r1 = p1.add_run(pg)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(10)
        if not item.startswith(" "):
            r1.bold = True

    doc.add_page_break()

    # ==========================================
    # PAGE 2: Abstract, Introduction, Objectives
    # ==========================================
    add_sec_heading("1. Abstract")
    add_body_p(
        "DUET MEAL is a lightweight, high-performance web and mobile meal management platform engineered specifically for the "
        "Dhaka University of Engineering & Technology (DUET) dining and canteen facility. Built entirely on a pure Core PHP RESTful API "
        "and a super lightweight Vanilla JavaScript single-page admin application, the system eliminates heavy framework overhead to deliver "
        "sub-second response times on shared local server hosting. In a bustling university environment, legacy manual coupon distribution and paper "
        "token registers suffer from record discrepancies, food wastage, and reconciliation delays. DUET MEAL resolves these challenges by providing a "
        "centralized, responsive digital hub where university staff (teachers and officers, both inside and outside campus residents) can book daily meals "
        "(Lunch and Dinner), reserve guest meals, and maintain a prepaid digital canteen wallet. With zero external framework dependencies on the backend, "
        "the system provides instantaneous wallet deductions, dynamic monthly settlement recalculations, and real-time administrative control."
    )

    add_sec_heading("2. Introduction")
    add_body_p(
        "Access to an organized, accountable dining facility is essential for university faculty and administrative staff. "
        "At Dhaka University of Engineering & Technology (DUET), hundreds of teachers and officers rely on the central canteen for daily catering. "
        "The traditional meal management process depended on physical token distribution, manual ledger books, and cash handovers at the counter. "
        "This led to billing inaccuracies, lack of advance catering demand visibility, and significant administrative workload during monthly reconciliations."
    )
    add_body_p(
        "DUET MEAL introduces a streamlined, ultra-lightweight software ecosystem to overcome these problems. The architecture is powered by a high-speed "
        "Core PHP backend serving clean RESTful JSON endpoints, paired with an asynchronous Vanilla JavaScript / Tailwind CSS single-page admin portal "
        "and an Android mobile application. Operating on a prepaid digital wallet model, users can schedule meals across flexible date ranges with strict "
        "cutoff enforcement (such as booking before 12:00 PM for lunch and 24 hours advance notice for guest dinners). For dining administrators, the platform "
        "delivers real-time headcount dashboards, instant cash recharge approvals, menu publishing, and automated monthly settlement adjustments without requiring "
        "heavy server compute."
    )

    add_sec_heading("3. Objectives")
    add_body_p(
        "The primary objective of this project is to develop a super lightweight, responsive, and secure digital meal booking and "
        "wallet management web application and RESTful API for the DUET dining community."
    )
    add_body_p("Specific Objectives:")
    add_bullet_p("To design a clean and responsive User Interface (UI)", "that delivers high speed and optimal rendering across desktop browsers, tablets, and mobile devices without heavy runtime bloat.")
    add_bullet_p("To build a high-performance Core PHP REST API", "handling user authentication, dining schedules, digital wallet ledger transactions, and automated business cutoffs efficiently.")
    add_bullet_p("To implement a reactive Vanilla JavaScript Admin Portal", "enabling canteen managers to control live booking toggles, audit cash collections, manage users, and approve wallet recharges asynchronously via AJAX/Fetch.")
    add_bullet_p("To develop an automated Meal Booking & Guest Management module", "supporting multi-day calendar scheduling, Lunch and Dinner selections, and guest reservations (up to 10 guests) with strict policy cutoffs.")
    add_bullet_p("To integrate an automated Monthly Settlement Engine", "that recalculates actual consumed meal rates at month-end and adjusts members' balances across the entire university seamlessly.")

    doc.add_page_break()

    # ==========================================
    # PAGE 3: Key Features & Methodology
    # ==========================================
    add_sec_heading("4. Key Features")
    add_body_p("The DUET MEAL platform is packed with high-efficiency features designed for both university dining members and administrators:")
    
    add_sec_heading("User Features:", level=2)
    add_bullet_p("Institutional Authentication", "Secure signup and login utilizing DUET email credentials (@duet.edu.bd) with encrypted password verification.")
    add_bullet_p("Smart Dashboard", "Real-time overview of current wallet balance, today's Lunch/Dinner status, menu items, and booking cutoff countdowns.")
    add_bullet_p("Date-Range Meal Booking", "Interactive calendar booking system allowing single or multi-day reservations with independent toggles for Lunch, Dinner, and guests.")
    add_bullet_p("Digital Canteen Wallet", "Prepaid balance tracking with a structured transaction ledger recording cash recharges (+), meal deductions (-), and monthly settlements.")
    add_bullet_p("Meal History & Consumption Tracking", "Filterable chronological records of all past bookings with live status indicators (Booked, Consumed, Auto-Cancelled).")
    add_bullet_p("Profile & Dining Preferences", "Management of user contact details, residential designation (Inside/Outside Campus), and meal reminder notifications.")
    add_bullet_p("FAQ & Bulletin Board", "Access to official dining announcements categorized by Dining, Academic, and Maintenance updates with unread badges.")

    add_sec_heading("Admin Features:", level=2)
    add_bullet_p("Interactive Dashboard", "Single-page administrative view of total users, today's confirmed Lunch/Dinner headcounts, pending recharges, and live booking master switch.")
    add_bullet_p("Asynchronous User Management", "Fast, client-rendered member management allowing administrators to search, inspect balances, and update user statuses without page reloads.")
    add_bullet_p("Menu & Schedule Management", "Intuitive tools to update daily menus, assign dining venues, configure meal timings, and enforce cutoff hours.")
    add_bullet_p("Cash Audit & Recharge Approvals", "Real-time recharge verification queue allowing managers to approve cash deposits and instantly credit member wallets.")
    add_bullet_p("Automated Monthly Settlement", "Single-click rate settlement engine that recalculates finalized per-meal costs (e.g., ৳88.10) and credits/debits balances automatically.")

    add_sec_heading("5. Methodology")
    add_bullet_p("Requirement Analysis", "Analyzed the specific needs of DUET dining members, evaluating the shortcomings of manual registers, cutoff violations, and monthly bill disputes.")
    add_bullet_p("System Design & Architecture", "Engineered a normalized MySQL schema and designed an ultra-lightweight architecture: zero-dependency Core PHP backend paired with a Vanilla JavaScript frontend.")
    add_bullet_p("Backend & REST API Development", "Developed a modular Core PHP RESTful API utilizing native PDO for parameterized database queries, JSON serialization, HTTP status code standards, and custom routing helpers.")
    add_bullet_p("Super Lightweight Frontend & Admin Portal", "Constructed the administrative single-page dashboard using modern Vanilla JavaScript (Fetch API, DOM manipulation, Lucide icons, and Tailwind CSS), achieving instantaneous page transitions without frontend build steps.")
    add_bullet_p("Functional Testing & Validation", "Conducted extensive testing on transactional wallet debiting, concurrent booking locks, 12:00 PM cutoff enforcement, and bulk monthly rate settlement math.")
    add_bullet_p("Local Server Deployment", "Deployed and validated the solution on a local XAMPP (Apache, PHP 8.x, MySQL) environment, guaranteeing smooth execution and low memory footprint.")

    doc.add_page_break()

    # ==========================================
    # PAGE 4: Project Overview - Auth
    # ==========================================
    add_sec_heading("6. Project Overview")
    
    add_sec_heading("6.0 User Login and Sign Up Section", level=2)
    add_body_p(
        "The User Login and Sign Up sections of the website were developed to provide university members with secure "
        "access and account creation functionality. The Sign Up section allowed new users to register by providing "
        "their full name, institutional email address, mobile number, campus residency category, and a password."
    )

    add_sec_heading("6.0.1 User Sign In Section", level=3)
    add_body_p(
        "The User Sign In section of the website was developed to provide registered users with secure access to "
        "their accounts. In this section, users were required to enter their university email address and password to log in. "
        "The Core PHP backend verifies the credentials against encrypted database records before granting access, ensuring that only authorized "
        "university personnel can access personal dining schedules, wallet balances, and booking history. This section played a crucial role in "
        "maintaining account security and providing a personalized experience on the DUET MEAL website."
    )
    add_screenshot_box("Figure-6.0.1 : User Sign In Page")

    add_sec_heading("6.0.2 User Sign Up Section", level=3)
    add_body_p(
        "The User Sign Up section was developed to allow university teachers and officers to create a new dining account. "
        "In this section, users enter their full name, official institutional email (@duet.edu.bd), contact number, and select their residential "
        "designation (Inside/Outside Campus). The system validates the input, initializes a dedicated prepaid digital wallet with a zero balance, "
        "and establishes account privileges securely."
    )
    add_screenshot_box("Figure-6.0.2 : User Sign Up Page")

    doc.add_page_break()

    # ==========================================
    # PAGE 5: Home, Booking, Guests
    # ==========================================
    add_sec_heading("6.1 Home Dashboard & Daily Meal Status", level=2)
    add_body_p(
        "The Home Dashboard section greets users with a clear overview of today's dining schedule. "
        "Powered by lightweight Core PHP API calls, it provides immediate visual indicators for booked meals, lunch and dinner cutoff countdown timers, "
        "current wallet balance, and quick links to make reservations, check notices, or view wallet history at a glance."
    )
    add_screenshot_box("Figure-6.1 : Home Dashboard & Daily Meal Status Page", height_in_inches=1.8)

    add_sec_heading("6.2 Meal Booking & Date-Range Selection", level=2)
    add_body_p(
        "The Meal Booking section allows users to plan and schedule their meals across specific dates. "
        "Users can select any date range on the interactive calendar and toggle Lunch, Dinner, or both. "
        "The system calculates the exact cost in real-time at the standard flat rate (৳90.00) and displays a transparent breakdown before confirmation."
    )
    add_screenshot_box("Figure-6.2 : Meal Booking & Date-Range Selection Page", height_in_inches=1.8)

    add_sec_heading("6.3 Check Availability & Guest Meal Reservation", level=2)
    add_body_p(
        "The Guest Meal reservation system enables university staff to book additional plates for visiting colleagues and guests (0 to 10 guests). "
        "The system automatically enforces university policies such as requiring guest dinner bookings to be placed at least 24 hours in advance, "
        "ensuring the canteen kitchen receives timely demand figures."
    )
    add_screenshot_box("Figure-6.3 : Check Availability & Guest Meal Reservation Dialog", height_in_inches=1.8)

    doc.add_page_break()

    # ==========================================
    # PAGE 6: Wallet & Daily Menu
    # ==========================================
    add_sec_heading("6.4 Review and Manage Your Canteen Wallet", level=2)
    add_body_p(
        "The Review and Manage Your Canteen Wallet section presents a clean overview of the user's financial standing. "
        "It displays the Total Deposited Balance and Available Spendable Balance. "
        "Users can initiate recharge requests and view detailed transaction logs showing cash credits, meal debits, and monthly settlement adjustments."
    )
    add_screenshot_box("Figure-6.4 : Review and Manage Your Canteen Wallet Page", height_in_inches=2.0)

    add_sec_heading("6.5 Daily Menu & Cutoff Schedule", level=2)
    add_body_p(
        "The Daily Menu & Cutoff Schedule section showcases the daily catering items for both Lunch and Dinner. "
        "Each meal card displays the venue (Main Canteen), service hours (12:30 PM – 2:00 PM for lunch; 7:30 PM – 9:00 PM for dinner), "
        "and the strict booking cutoff time (12:00 PM for lunch), ensuring users never miss booking deadlines."
    )
    add_screenshot_box("Figure-6.5 : Daily Menu & Cutoff Schedule View", height_in_inches=2.0)

    doc.add_page_break()

    # ==========================================
    # PAGE 7: History & Profile
    # ==========================================
    add_sec_heading("6.6 My Order & Meal History", level=2)
    add_body_p(
        "The My Order & Meal History section of the website was implemented to allow users to view their complete meal booking "
        "history in an organized manner. This section displayed important information such as booking ID, date, "
        "ordered meals (Lunch/Dinner), guest count, total cost, and current status (Booked, Consumed, or Auto-Cancelled). "
        "The My Order section improved transparency and user satisfaction by keeping users informed about their dining history."
    )
    add_screenshot_box("Figure-6.6 : My Order & Meal History Page", height_in_inches=2.1)

    add_sec_heading("6.7 My Profile Section", level=2)
    add_body_p(
        "The My Profile section of the website was designed to provide users with a personalized area to view "
        "and manage their personal information. This section displayed user details such as name, institutional email address, "
        "mobile number, campus designation, and residential status. Users were allowed to update selected preferences, "
        "including phone number and meal reminder toggles, while institutional credentials remained secured."
    )
    add_screenshot_box("Figure-6.7 : My Profile Section Page", height_in_inches=2.1)

    doc.add_page_break()

    # ==========================================
    # PAGE 8: FAQ & Notice Board
    # ==========================================
    add_sec_heading("6.8 FAQ & Notice Board Section", level=2)
    add_body_p(
        "The FAQ & Notice Board section of the website was developed to provide quick answers to common user queries "
        "related to meal booking cutoffs, wallet recharge verification, guest policies, and monthly settlements. "
        "This section presented a structured list of frequently asked questions along with their corresponding answers, "
        "reducing user confusion and minimizing direct administrative inquiries."
    )
    add_screenshot_box("Figure-6.8 : Frequently Asked Questions (FAQ) Section", height_in_inches=2.1)

    add_body_p(
        "In addition, the Notice Board subsection delivers real-time official bulletins from the canteen committee. "
        "Notices are categorized into Dining, Academic, and Maintenance updates, complete with unread indicator badges "
        "to ensure university members are promptly notified of policy updates or rate finalizations."
    )
    add_screenshot_box("Figure-6.8.1 : University Notice Board & Bulletins View", height_in_inches=2.1)

    doc.add_page_break()

    # ==========================================
    # PAGE 9: Admin Panel - Sign In
    # ==========================================
    add_sec_heading("7. Admin Panel")
    add_body_p(
        "The Admin Panel is a super lightweight Single-Page Application (SPA) built using pure Vanilla JavaScript, "
        "asynchronous Fetch API calls, and Tailwind CSS. It interacts with the high-speed Core PHP REST API to give "
        "canteen managers total operational control over daily meal headcounts, cash audit records, user accounts, "
        "catering menus, and month-end rate settlements with zero page reloads."
    )

    add_sec_heading("7.1 Admin Sign-In Section", level=2)
    add_body_p(
        "The Admin Sign-In section of the website was designed to provide a secure access point for the administrator. "
        "In this section, the admin was required to enter a valid username/email and password to gain access to the admin panel. "
        "Proper authentication and validation mechanisms were implemented to prevent unauthorized access, ensuring that only authorized "
        "canteen personnel could manage users, recharge wallets, and update meal configurations."
    )
    add_screenshot_box("Figure-7.1 : Admin Sign In Page", height_in_inches=3.0)

    doc.add_page_break()

    # ==========================================
    # PAGE 10: Admin Dashboard & User Management
    # ==========================================
    add_sec_heading("7.2 Admin Dashboard Section", level=2)
    add_body_p(
        "The Admin Dashboard section of the website was developed to provide an overview of the system’s "
        "current status. This section displayed important statistical information such as total registered users, "
        "today's Lunch & Dinner bookings, pending recharge requests, and active ledger balance. It also features a Master Booking Switch "
        "to toggle the live booking service in real time, allowing administrators to monitor canteen operations efficiently."
    )
    add_screenshot_box("Figure-7.2 : Admin Dashboard Section Overview", height_in_inches=2.1)

    add_sec_heading("7.3 Admin User Management Section", level=2)
    add_body_p(
        "The Admin User Management section of the website was designed to allow the administrator to view "
        "and manage registered users. In this section, the administrator could access detailed information about "
        "each user, including their name, institutional email, designation (Teacher/Officer), campus residential status (Inside/Outside), "
        "and active wallet balance. Fast client-side searching and status management maintain complete database integrity."
    )
    add_screenshot_box("Figure-7.3 : Admin User Management Section", height_in_inches=2.1)

    doc.add_page_break()

    # ==========================================
    # PAGE 11: Admin Menu & Order Management
    # ==========================================
    add_sec_heading("7.4 Admin Menu Management Section", level=2)
    add_body_p(
        "The Admin Menu Management section was developed to enable canteen managers to schedule daily catering meals. "
        "In this section, the administrator could input lunch and dinner items, assign dining halls, set service timings, "
        "and configure cutoff thresholds. Dynamic JavaScript updates communicate directly with the Core PHP API, ensuring "
        "kitchen staff receive precise demand headcounts instantly."
    )
    add_screenshot_box("Figure-7.4 : Admin Menu Management Section", height_in_inches=2.1)

    add_sec_heading("7.5 Admin Recharge Approvals & Transactions", level=2)
    add_body_p(
        "The Admin Recharge Approvals & Transactions section was developed to allow the administrator to monitor "
        "and approve financial transactions. The administrator could review cash deposit requests from faculty, "
        "verify amounts, and click 'Approve' to instantly credit digital wallets asynchronously. Detailed immutable transaction "
        "logs ensure complete financial auditing compliance."
    )
    add_screenshot_box("Figure-7.5 : Admin Recharge Approvals & Transactions", height_in_inches=2.1)

    doc.add_page_break()

    # ==========================================
    # PAGE 12: Admin Settlement & Footer
    # ==========================================
    add_sec_heading("7.6 Admin Monthly Settlement Section", level=2)
    add_body_p(
        "The Admin Monthly Settlement section was designed to automate month-end financial reconciliations. "
        "At the end of each billing cycle, administrators input the finalized actual meal rate (e.g., ৳88.10). "
        "The Core PHP backend recalculates charges across all consumed meals and applies automated wallet adjustments across all member accounts "
        "with sub-second execution speed."
    )
    add_screenshot_box("Figure-7.6 : Admin Monthly Settlement Section", height_in_inches=1.8)

    add_sec_heading("8. Footer", level=1)
    add_body_p(
        "The Footer section of the website was developed to provide supplementary information and resources "
        "to users at the bottom of each page. It contained elements such as dining office contact details, quick navigation "
        "links, operating hours, and copyright information. The footer improved website credibility and usability."
    )

    add_sec_heading("8.1 Quick Links, Help, and Contact in Footer Section", level=2)
    add_body_p(
        "The Footer section of the website included Quick Links, Help, and Contact information to improve "
        "navigation and user support. The Quick Links provided direct access to important pages such as "
        "Home, Daily Menu, Booking, Canteen Wallet, and Notices. The Contact information displayed canteen office phone numbers "
        "and email addresses, enabling users to reach administration easily."
    )
    add_screenshot_box("Figure-8.1 : Quick Links, Help, and Contact in Footer Section", height_in_inches=1.8)

    doc.add_page_break()

    # ==========================================
    # PAGE 13: Challenges, Limitations, Outcomes, Conclusion
    # ==========================================
    add_sec_heading("9. Challenges")
    add_body_p(
        "During the development of the DUET MEAL website and API, several technical challenges were addressed: "
        "Building a high-throughput REST API strictly in pure Core PHP without relying on heavy frameworks required creating robust custom "
        "routing, PDO database wrappers, JSON response handlers, and CORS headers from scratch. Implementing a reactive single-page admin panel "
        "in pure Vanilla JavaScript required careful state handling, asynchronous DOM updates via Fetch API, and toast notification queues. "
        "Additionally, enforcing precise time-based cutoff policies (12:00 PM lunch cutoff and 24-hour advance guest dinner rule) and managing "
        "concurrency-locked wallet transactions to prevent double-spending during peak booking hours required rigorous database transaction handling."
    )

    add_sec_heading("10. Limitations")
    add_body_p(
        "The DUET MEAL website had some limitations. Direct online payment gateway integration (such as bKash, Nagad, or bank cards) "
        "was not included in this version due to university merchant account onboarding timelines, relying instead on admin-verified cash recharges. "
        "Hardware-level turnstile RFID/biometric scanning at the dining hall entrance is currently simulated via QR token verification, "
        "with dedicated IoT hardware integration planned for subsequent releases."
    )

    add_sec_heading("11. Learning Outcomes")
    add_body_p("From this project, several key learning outcomes were achieved:")
    add_bullet_p("Core PHP REST API Development", "Gained deep experience in architecting lightweight, high-performance RESTful APIs from scratch using native PHP and PDO without external framework bloat.")
    add_bullet_p("Vanilla JavaScript SPA Architecture", "Mastered reactive single-page frontend development using modern JavaScript, Fetch API, DOM manipulation, and asynchronous workflows.")
    add_bullet_p("Fintech & Prepaid Wallet Engineering", "Developed skills in managing prepaid digital wallets, transaction locking (ACID compliance), and automated bulk monthly settlements.")
    add_bullet_p("Secure Role-Based Authentication", "Learned to implement institutional role-based authentication and secure session management for university staff.")
    add_bullet_p("UI/UX & Responsive Design", "Created clean, intuitive, and mobile-friendly interfaces optimized for low-latency university server environments.")

    add_sec_heading("12. Conclusion")
    add_body_p(
        "In conclusion, the DUET MEAL platform was successfully designed and implemented as a super lightweight, high-performance "
        "canteen management and meal booking system. The project demonstrated the effective integration of a pure Core PHP RESTful API backend "
        "and a reactive Vanilla JavaScript admin portal. Key capabilities including date-range meal scheduling, guest reservations, digital prepaid "
        "wallet management, cash audit approvals, and automated monthly settlements operate seamlessly with minimal server resource consumption. "
        "Overall, this project provided valuable insights into creating robust, scalable, and ultra-fast institutional software solutions."
    )

    output_path = r"d:\xampp\htdocs\duetmealapi\DUET_Meal_Project_Report.docx"
    doc.save(output_path)
    print(f"Report generated successfully at: {output_path}")

if __name__ == '__main__':
    create_report()
